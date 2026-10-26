import csv
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import load_workbook
from app.services import generate_excel as report
from app.services import garmin_sync


def activity(year, name, distance=10000):
    return SimpleNamespace(
        id=year, garmin_activity_id=year, name=name, sport_type="running",
        start_date=datetime(year, 12, 31, 10), raw_summary_json={},
        distance_m=distance, moving_time_s=3600, duration_s=3600,
        elapsed_time_s=3600, elevation_gain_m=100, elevation_loss_m=90,
        average_heartrate=140, max_heartrate=160, average_cadence_spm=170,
        average_power_w=200, normalized_power_w=210, max_power_w=300,
        aerobic_training_effect=2, anaerobic_training_effect=1, vo2max=50,
        calories=600,
    )


class ExportTests(unittest.TestCase):
    def test_annual_global_csv_and_selective_regeneration(self):
        items = [activity(2023, "Ancienne"), activity(2024, "Nouvelle")]
        with tempfile.TemporaryDirectory() as directory, \
             patch.object(report, "load_activities", return_value=items), \
             patch.object(report, "route_record_rows", return_value=[]):
            root = Path(directory)
            report.generate_exports(export_path=root / "Sport_report.xlsx")
            annual = load_workbook(root / "Sport_report_2023.xlsx")
            self.assertEqual(annual.sheetnames, ["Dashboard", "Semaines", "Activites", "Records", "_Data"])
            self.assertEqual(annual["Activites"]["C4"].value, "Ancienne")
            self.assertEqual(annual["Dashboard"]["B9"].value, 10)
            self.assertIsNotNone(annual["_Data"]["H13"].value)
            global_book = load_workbook(root / "Sport_report.xlsx")
            self.assertEqual(global_book.sheetnames, ["All", "Records"])
            with (root / "Sport_activities.csv").open(encoding="utf-8-sig", newline="") as file:
                rows = list(csv.reader(file, delimiter=";"))
            self.assertEqual(len(rows), 3)
            self.assertEqual(rows[1][2], "Nouvelle")
            self.assertEqual(rows[1][6], "3600.0")
            old = (root / "Sport_report_2023.xlsx").read_bytes()
            report.generate_exports([2024], export_path=root / "Sport_report.xlsx")
            self.assertEqual((root / "Sport_report_2023.xlsx").read_bytes(), old)

    def test_sync_tracks_all_new_local_years(self):
        activities = [
            {"startTimeLocal": "2024-12-31 12:00:00", "startTimeGMT": "2024-12-31 11:00:00"},
            {"startTimeLocal": "2025-01-01 00:15:00", "startTimeGMT": "2024-12-31 23:15:00"},
            {"startTimeLocal": "2023-01-01 12:00:00", "startTimeGMT": "2023-01-01 11:00:00"},
        ]
        api = SimpleNamespace(get_activities=lambda offset, limit: activities)
        with patch.object(garmin_sync, "sync_activity", side_effect=[True, True, False]), \
             patch.object(garmin_sync.time, "sleep"):
            count, years, undated = garmin_sync.get_history(api, max_activities=3)
        self.assertEqual(count, 2)
        self.assertEqual(years, {2024, 2025})
        self.assertFalse(undated)

    def test_local_year_at_new_year_boundary(self):
        item = activity(2024, "Minuit")
        item.raw_summary_json = {"startTimeLocal": "2025-01-01 00:15:00"}
        self.assertEqual(report.activity_day(item).year, 2025)


if __name__ == "__main__":
    unittest.main()
