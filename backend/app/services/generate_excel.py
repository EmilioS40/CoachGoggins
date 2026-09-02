import math
from collections import defaultdict
from copy import copy
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.axis import ChartLines, DateAxis
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import to_excel
from openpyxl.worksheet.table import Table, TableStyleInfo
from sqlalchemy import select

from app.db.session import SessionLocal
from app.models.activity import Activity
from app.models.activity_record import ActivityRecord


EXPORT_PATH = Path("/var/lib/sport/exports/sport.xlsx")

RUNNING_COLOR = "60A5FA"
TRAIL_COLOR = "1D4ED8"
WALK_COLOR = "22C55E"
BIKE_COLOR = "F97316"
OTHER_COLOR = "9CA3AF"

CHART_BLUE = "2563EB"
CHART_GREEN = "15803D"
CHART_ORANGE = "EA580C"
CHART_TEAL = "0F766E"

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(size=18, bold=True, color="111827")
SUBTITLE_FONT = Font(size=11, color="6B7280")
SECTION_FONT = Font(size=12, bold=True, color="111827")
CARD_FILL = PatternFill("solid", fgColor="F3F4F6")
ACCENT_FILL = PatternFill("solid", fgColor="DBEAFE")
MAXIMUM_FILL = PatternFill("solid", fgColor="FEF08A")
THIN = Side(style="thin", color="D1D5DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def safe_round(value, digits=1):
    if value is None:
        return None
    return round(float(value), digits)


def duration_days(seconds):
    if not seconds:
        return None
    return float(seconds) / 86400


def pace_days(duration_s, distance_m):
    if not duration_s or not distance_m:
        return None
    km = distance_m / 1000
    if km <= 0:
        return None
    return (duration_s / km) / 86400


def duration_seconds(activity):
    return activity.moving_time_s or activity.duration_s or activity.elapsed_time_s or 0


def parse_local_datetime(activity):
    raw = activity.raw_summary_json or {}
    value = raw.get("startTimeLocal")

    if value:
        try:
            return datetime.strptime(value, "%Y-%m-%d %H:%M:%S")
        except ValueError:
            pass

    if activity.start_date:
        return activity.start_date.replace(tzinfo=None)

    return None


def activity_day(activity):
    local_dt = parse_local_datetime(activity)
    return local_dt.date() if local_dt else None


def activity_type_text(activity):
    text = " ".join(
        str(part or "")
        for part in [
            activity.sport_type,
            activity.name,
        ]
    ).lower()
    return (
        text.replace("é", "e")
        .replace("è", "e")
        .replace("ê", "e")
        .replace("à", "a")
        .replace("ù", "u")
    )


def sport_category(activity):
    text = activity_type_text(activity)

    if "trail" in text:
        return "trail"
    if "run" in text or "course" in text:
        return "running"
    if "walk" in text or "marche" in text or "hike" in text or "rando" in text:
        return "walk"
    if "bike" in text or "cycling" in text or "cycl" in text or "velo" in text:
        return "bike"
    return "other"


def is_running_or_trail(activity):
    return sport_category(activity) in {"running", "trail"}


def is_route_running(activity):
    return sport_category(activity) == "running"


def week_start(day):
    return day - timedelta(days=day.weekday())


def iso_week_label(day):
    iso = day.isocalendar()
    return f"{iso.year}-S{iso.week:02d}"


def latest_non_null(activities, attr):
    for activity in sorted(
        activities,
        key=lambda item: parse_local_datetime(item) or datetime.min,
        reverse=True,
    ):
        value = getattr(activity, attr)
        if value is not None:
            return value
    return None


def load_activities():
    db = SessionLocal()
    try:
        return db.scalars(select(Activity).order_by(Activity.start_date.asc())).all()
    finally:
        db.close()


def activity_rows(activities):
    rows = []
    for activity in sorted(
        activities,
        key=lambda item: parse_local_datetime(item) or datetime.min,
        reverse=True,
    ):
        local_dt = parse_local_datetime(activity)
        duration_s = duration_seconds(activity)

        rows.append([
            local_dt.date() if local_dt else None,
            local_dt.time() if local_dt else None,
            activity.name,
            sport_category(activity),
            activity.sport_type,
            safe_round((activity.distance_m or 0) / 1000, 2),
            duration_days(duration_s),
            pace_days(duration_s, activity.distance_m),
            safe_round(activity.elevation_gain_m, 0),
            safe_round(activity.elevation_loss_m, 0),
            safe_round(activity.average_heartrate, 0),
            safe_round(activity.max_heartrate, 0),
            safe_round(activity.average_cadence_spm, 1),
            safe_round(activity.average_power_w, 0),
            safe_round(activity.normalized_power_w, 0),
            safe_round(activity.max_power_w, 0),
            safe_round(activity.aerobic_training_effect, 1),
            safe_round(activity.anaerobic_training_effect, 1),
            safe_round(activity.vo2max, 1),
            safe_round(activity.calories, 0),
            activity.garmin_activity_id,
        ])

    return rows


def weekly_rows(activities):
    buckets = defaultdict(list)

    for activity in activities:
        if not is_running_or_trail(activity):
            continue

        day = activity_day(activity)
        if day:
            buckets[week_start(day)].append(activity)

    rows = []

    for start_day in sorted(buckets, reverse=True):
        items = buckets[start_day]
        total_distance_m = sum(a.distance_m or 0 for a in items)
        total_duration_s = sum(duration_seconds(a) for a in items)
        total_gain = sum(a.elevation_gain_m or 0 for a in items)
        total_loss = sum(a.elevation_loss_m or 0 for a in items)
        longest = max((a.distance_m or 0 for a in items), default=0)
        max_gain = max((a.elevation_gain_m or 0 for a in items), default=0)
        running_count = sum(1 for a in items if sport_category(a) == "running")
        trail_count = sum(1 for a in items if sport_category(a) == "trail")

        hr_weight = sum(
            (a.average_heartrate or 0) * duration_seconds(a)
            for a in items
            if a.average_heartrate and duration_seconds(a)
        )
        power_weight = sum(
            (a.average_power_w or 0) * duration_seconds(a)
            for a in items
            if a.average_power_w and duration_seconds(a)
        )
        weighted_duration_hr = sum(
            duration_seconds(a) for a in items if a.average_heartrate
        )
        weighted_duration_power = sum(
            duration_seconds(a) for a in items if a.average_power_w
        )

        rows.append([
            start_day.isocalendar().year,
            start_day.isocalendar().week,
            start_day,
            iso_week_label(start_day),
            len(items),
            running_count,
            trail_count,
            round(total_distance_m / 1000, 2),
            total_duration_s / 3600,
            pace_days(total_duration_s, total_distance_m),
            round(total_gain, 0),
            round(total_loss, 0),
            round((total_distance_m / 1000) / len(items), 2) if items else None,
            round(longest / 1000, 2),
            round(max_gain, 0),
            round(hr_weight / weighted_duration_hr, 0) if weighted_duration_hr else None,
            round(power_weight / weighted_duration_power, 0) if weighted_duration_power else None,
        ])

    return rows


def period_stats(activities, start_day, end_day):
    items = [
        activity
        for activity in activities
        if is_running_or_trail(activity)
        and activity_day(activity)
        and start_day <= activity_day(activity) < end_day
    ]

    distance = sum(a.distance_m or 0 for a in items)
    duration = sum(duration_seconds(a) for a in items)
    gain = sum(a.elevation_gain_m or 0 for a in items)

    return {
        "activities": len(items),
        "running": sum(1 for a in items if sport_category(a) == "running"),
        "trail": sum(1 for a in items if sport_category(a) == "trail"),
        "distance_km": round(distance / 1000, 1),
        "duration_h": round(duration / 3600, 1),
        "elevation_gain_m": round(gain, 0),
        "pace": pace_days(duration, distance),
    }


def current_periods(today):
    week = week_start(today)
    month = today.replace(day=1)

    if month.month == 12:
        next_month = date(month.year + 1, 1, 1)
    else:
        next_month = date(month.year, month.month + 1, 1)

    return [
        ("Semaine en cours", week, week + timedelta(days=7)),
        ("Mois en cours", month, next_month),
        ("Annee en cours", date(today.year, 1, 1), date(today.year + 1, 1, 1)),
    ]


def daily_year_metric_data(activities, year, today, value_func):
    totals = defaultdict(float)

    for activity in activities:
        day = activity_day(activity)
        if day and day.year == year and is_running_or_trail(activity):
            totals[day] += value_func(activity)

    rows = []
    cumulative = 0.0
    day = date(year, 1, 1)

    end = date(year, 12, 31)

    while day <= end:
        if day <= today:
            cumulative += totals[day]
            value = round(cumulative, 2)
        else:
            value = None

        rows.append([day, value])
        day += timedelta(days=1)

    return rows


def weekly_year_metric_data(activities, year, today, value_func):
    totals = defaultdict(float)

    for activity in activities:
        day = activity_day(activity)
        if day and day.year == year and is_running_or_trail(activity):
            totals[week_start(day)] += value_func(activity)

    rows = []
    current = week_start(date(year, 1, 1))
    current_week = week_start(today)
    end = week_start(date(year, 12, 31))

    while current <= end:
        value = round(totals[current], 2) if current <= current_week else None
        rows.append([current, value])
        current += timedelta(days=7)

    return rows


def daily_year_data(activities, year, today):
    return daily_year_metric_data(
        activities,
        year,
        today,
        lambda activity: (activity.distance_m or 0) / 1000,
    )


def daily_year_elevation_data(activities, year, today):
    return daily_year_metric_data(
        activities,
        year,
        today,
        lambda activity: activity.elevation_gain_m or 0,
    )


def weekly_year_data(activities, year, today):
    return weekly_year_metric_data(
        activities,
        year,
        today,
        lambda activity: (activity.distance_m or 0) / 1000,
    )


def weekly_year_elevation_data(activities, year, today):
    return weekly_year_metric_data(
        activities,
        year,
        today,
        lambda activity: activity.elevation_gain_m or 0,
    )


def monthly_year_data(activities, year, today):
    totals = defaultdict(float)

    for activity in activities:
        day = activity_day(activity)
        if day and day.year == year and is_running_or_trail(activity):
            totals[(day.year, day.month)] += (activity.distance_m or 0) / 1000

    rows = []

    for month in range(1, 13):
        value = round(totals[(year, month)], 2) if month <= today.month else None
        rows.append([date(year, month, 1), value])

    return rows


def yearly_rows(activities):
    buckets = defaultdict(list)

    for activity in activities:
        day = activity_day(activity)
        if day and is_running_or_trail(activity):
            buckets[day.year].append(activity)

    rows = []

    for year in sorted(buckets):
        items = buckets[year]
        total_distance_m = sum(a.distance_m or 0 for a in items)
        total_duration_s = sum(duration_seconds(a) for a in items)
        total_gain = sum(a.elevation_gain_m or 0 for a in items)
        total_loss = sum(a.elevation_loss_m or 0 for a in items)
        longest = max((a.distance_m or 0 for a in items), default=0)
        running_count = sum(1 for a in items if sport_category(a) == "running")
        trail_count = sum(1 for a in items if sport_category(a) == "trail")

        hr_weight = sum(
            (a.average_heartrate or 0) * duration_seconds(a)
            for a in items
            if a.average_heartrate and duration_seconds(a)
        )
        weighted_duration_hr = sum(
            duration_seconds(a) for a in items if a.average_heartrate
        )

        rows.append([
            year,
            len(items),
            running_count,
            trail_count,
            round(total_distance_m / 1000, 2),
            round(total_duration_s / 3600, 1),
            round(total_gain, 0),
            round(total_loss, 0),
            round(longest / 1000, 2),
            round(hr_weight / weighted_duration_hr, 0) if weighted_duration_hr else None,
            safe_round(max((a.max_power_w or 0 for a in items), default=0), 0),
            safe_round(max((a.vo2max or 0 for a in items), default=0), 1),
        ])

    return rows


def calendar_day_categories(activities, year):
    priority = {
        "trail": 4,
        "running": 3,
        "bike": 2,
        "walk": 1,
        "other": 0,
    }
    days = {}

    for activity in activities:
        day = activity_day(activity)
        if not day or day.year != year:
            continue

        category = sport_category(activity)

        if day not in days or priority[category] > priority[days[day]]:
            days[day] = category

    return days


def yearly_day_breakdown(activities, year, today):
    category_by_day = calendar_day_categories(activities, year)
    labels = [
        ("Course a pied", "running", RUNNING_COLOR),
        ("Trail", "trail", TRAIL_COLOR),
        ("Velo", "bike", BIKE_COLOR),
        ("Marche", "walk", WALK_COLOR),
        ("Autre", "other", OTHER_COLOR),
        ("Sans activite", "inactive", "E5E7EB"),
    ]
    counts = defaultdict(int)

    current = date(year, 1, 1)
    while current <= today:
        counts[category_by_day.get(current, "inactive")] += 1
        current += timedelta(days=1)

    total_days = sum(counts.values())
    rows = []

    for label, category, color in labels:
        count = counts[category]
        rows.append((label, count, count / total_days if total_days else 0, color))

    return rows, total_days


def fastest_split(records, target_m):
    points = [
        (record.distance_m, record.timestamp)
        for record in records
        if record.distance_m is not None and record.timestamp is not None
    ]

    best = None
    left = 0

    for right, (right_distance, right_time) in enumerate(points):
        while left < right and right_distance - points[left][0] >= target_m:
            elapsed = (right_time - points[left][1]).total_seconds()

            if elapsed > 0 and (best is None or elapsed < best[0]):
                best = (elapsed, points[left][0], right_distance)

            left += 1

    return best


def route_record_rows(activities):
    route_activities = [
        activity
        for activity in activities
        if is_route_running(activity) and (activity.distance_m or 0) >= 1000
    ]

    targets = [
        ("5 km", 5000),
        ("10 km", 10000),
        ("Semi-marathon", 21097.5),
        ("Marathon", 42195),
    ]

    best = {label: None for label, _ in targets}
    db = SessionLocal()

    try:
        for activity in route_activities:
            activity_targets = [
                (label, target)
                for label, target in targets
                if (activity.distance_m or 0) >= target
            ]

            if not activity_targets:
                continue

            records = db.scalars(
                select(ActivityRecord)
                .where(ActivityRecord.activity_id == activity.id)
                .order_by(ActivityRecord.record_index.asc())
            ).all()

            for label, target in activity_targets:
                split = fastest_split(records, target)

                if not split:
                    continue

                elapsed = split[0]

                if best[label] is None or elapsed < best[label]["elapsed"]:
                    best[label] = {
                        "elapsed": elapsed,
                        "target": target,
                        "activity": activity,
                    }

    finally:
        db.close()

    rows = []

    for label, target in targets:
        item = best[label]

        if not item:
            rows.append([label, None, None, None, None, None])
            continue

        activity = item["activity"]
        elapsed = item["elapsed"]

        rows.append([
            label,
            duration_days(elapsed),
            pace_days(elapsed, target),
            activity_day(activity),
            activity.name,
            activity.garmin_activity_id,
        ])

    return rows


def trail_record_rows(activities):
    trails = [activity for activity in activities if sport_category(activity) == "trail"]

    def best_activity(label, value_func, unit):
        candidates = [activity for activity in trails if value_func(activity)]

        if not candidates:
            return [label, None, unit, None, None, None]

        activity = max(candidates, key=value_func)
        value = value_func(activity)

        return [
            label,
            round(value, 2),
            unit,
            activity_day(activity),
            activity.name,
            activity.garmin_activity_id,
        ]

    return [
        best_activity("Plus longue distance trail", lambda a: (a.distance_m or 0) / 1000, "km"),
        best_activity("Plus longue duree trail", lambda a: duration_seconds(a) / 3600, "h"),
        best_activity("Plus gros D+ trail", lambda a: a.elevation_gain_m or 0, "m"),
    ]


def general_record_rows(activities, weeks):
    sport_activities = [a for a in activities if is_running_or_trail(a)]

    def best_activity(label, value_func, unit):
        candidates = [activity for activity in sport_activities if value_func(activity)]

        if not candidates:
            return [label, None, unit, None, None, None]

        activity = max(candidates, key=value_func)
        value = value_func(activity)

        return [
            label,
            round(value, 2),
            unit,
            activity_day(activity),
            activity.name,
            activity.garmin_activity_id,
        ]

    def best_week(label, index, unit):
        if not weeks:
            return [label, None, unit, None, None, None]

        row = max(weeks, key=lambda item: item[index] or 0)
        return [label, row[index], unit, row[2], row[3], None]

    return [
        best_week("Meilleure semaine distance", 7, "km"),
        best_week("Meilleure semaine duree", 8, "h"),
        best_week("Meilleure semaine D+", 10, "m"),
        best_activity("Plus longue distance", lambda a: (a.distance_m or 0) / 1000, "km"),
        best_activity("Plus longue duree", lambda a: duration_seconds(a) / 3600, "h"),
        best_activity("Plus gros D+", lambda a: a.elevation_gain_m or 0, "m"),
        best_activity("FC max observee", lambda a: a.max_heartrate or 0, "bpm"),
        best_activity("Puissance max observee", lambda a: a.max_power_w or 0, "W"),
        best_activity("VO2max max", lambda a: a.vo2max or 0, ""),
    ]


def setup_sheet(ws, freeze="A2"):
    ws.freeze_panes = freeze
    ws.sheet_view.showGridLines = False


def write_title(ws, title, subtitle=None):
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT

    if subtitle:
        ws["A2"] = subtitle
        ws["A2"].font = SUBTITLE_FONT


def style_header(row):
    for cell in row:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        cell.border = BORDER


def style_section_cell(cell):
    cell.font = SECTION_FONT
    cell.fill = ACCENT_FILL
    cell.border = BORDER


def add_table(ws, name, first_row, last_row, last_col):
    if last_row <= first_row:
        return

    ref = f"A{first_row}:{get_column_letter(last_col)}{last_row}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def autosize(ws, max_width=34):
    for column in ws.columns:
        letter = column[0].column_letter
        width = 10

        for cell in column:
            value = cell.value

            if value is not None:
                width = max(width, min(max_width, len(str(value)) + 2))

        ws.column_dimensions[letter].width = width


def format_columns(ws, formats, start_row=4):
    for col, fmt in formats.items():
        for cell in ws[col][start_row - 1:]:
            cell.number_format = fmt


def make_table_sheet(ws, title, headers, rows, table_name, formatter=None):
    setup_sheet(ws)
    write_title(ws, title)
    ws.append([])
    ws.append(headers)
    style_header(ws[3])

    for row in rows:
        ws.append(row)

    if rows:
        add_table(ws, table_name, 3, ws.max_row, len(headers))

    if formatter:
        formatter(ws)

    autosize(ws)


def write_chart_data(
    ws,
    daily_rows,
    weekly_rows_current_year,
    monthly_rows_current_year,
    daily_elevation_rows,
    weekly_elevation_rows,
):
    ws.title = "_Data"
    ws.sheet_state = "hidden"

    ws["A1"] = "Date"
    ws["B1"] = "Cumul km"

    for index, row in enumerate(daily_rows, start=2):
        ws.cell(index, 1, row[0])
        ws.cell(index, 2, row[1])

    ws["D1"] = "Semaine"
    ws["E1"] = "Distance km"

    for index, row in enumerate(weekly_rows_current_year, start=2):
        ws.cell(index, 4, row[0])
        ws.cell(index, 5, row[1])
        ws.cell(index, 4).number_format = "dd mmm"

    ws["G1"] = "Mois"
    ws["H1"] = "Distance km"

    for index, row in enumerate(monthly_rows_current_year, start=2):
        ws.cell(index, 7, row[0])
        ws.cell(index, 8, row[1])
        ws.cell(index, 7).number_format = "mmm yyyy"

    ws["J1"] = "Date"
    ws["K1"] = "Cumul D+"

    for index, row in enumerate(daily_elevation_rows, start=2):
        ws.cell(index, 10, row[0])
        ws.cell(index, 11, row[1])

    ws["M1"] = "Semaine"
    ws["N1"] = "D+"

    for index, row in enumerate(weekly_elevation_rows, start=2):
        ws.cell(index, 13, row[0])
        ws.cell(index, 14, row[1])
        ws.cell(index, 13).number_format = "dd mmm"


def nice_axis_scale(values):
    maximum = max((float(value or 0) for value in values), default=0)

    if maximum <= 0:
        return 1, 0.2

    raw_step = maximum / 5
    magnitude = 10 ** math.floor(math.log10(raw_step))
    normalized = raw_step / magnitude

    if normalized <= 1:
        step = magnitude
    elif normalized <= 2:
        step = 2 * magnitude
    elif normalized <= 5:
        step = 5 * magnitude
    else:
        step = 10 * magnitude

    return math.ceil(maximum / step) * step, step


def axis_scale(values, major_unit=None):
    if major_unit is None:
        return nice_axis_scale(values)

    maximum = max((float(value or 0) for value in values), default=0)
    return max(major_unit, math.ceil(maximum / major_unit) * major_unit), major_unit


def style_line_series(series, color, marker_size=5):
    series.graphicalProperties.line.solidFill = color
    series.graphicalProperties.line.width = 25000
    series.marker.symbol = "circle"
    # Excel/OpenXML only accepts marker sizes from 2 to 72.
    series.marker.size = max(2, marker_size)
    series.marker.graphicalProperties.solidFill = color
    series.marker.graphicalProperties.line.solidFill = color
    series.smooth = False


def style_bar_series(series, color):
    series.graphicalProperties.solidFill = color
    series.graphicalProperties.line.solidFill = color
    # Match the clearer 5 pt outline used in Excel's series formatting.
    series.graphicalProperties.line.width = 63500


def add_line_chart(
    ws,
    data_ws,
    title,
    min_col,
    max_row,
    category_col,
    anchor,
    values,
    x_axis_title,
    y_axis_title,
    date_format,
    major_time_unit,
    color,
    y_major_unit=None,
    marker_size=5,
    x_min_date=None,
    x_max_date=None,
    height=9,
    width=18.5,
):
    chart = LineChart()
    chart.title = title
    chart.height = height
    chart.width = width

    chart.x_axis = DateAxis(axId=500, crossAx=100)
    chart.x_axis.number_format = date_format
    chart.x_axis.majorTimeUnit = major_time_unit
    chart.x_axis.tickLblPos = "nextTo"
    chart.x_axis.tickLblSkip = 1
    chart.x_axis.tickMarkSkip = 1
    chart.x_axis.tickMark = "out"
    chart.x_axis.delete = False

    if x_min_date is not None:
        chart.x_axis.scaling.min = to_excel(x_min_date)
    if x_max_date is not None:
        chart.x_axis.scaling.max = to_excel(x_max_date)

    chart.y_axis.axId = 100
    chart.y_axis.crossAx = 500
    chart.y_axis.numFmt = "#,##0"
    chart.y_axis.tickLblPos = "nextTo"
    chart.y_axis.tickMark = "out"
    chart.y_axis.delete = False
    chart.y_axis.scaling.min = 0
    chart.y_axis.scaling.max, chart.y_axis.majorUnit = axis_scale(
        values,
        y_major_unit,
    )
    chart.y_axis.majorGridlines = ChartLines()

    data = Reference(data_ws, min_col=min_col, min_row=1, max_row=max_row)
    cats = Reference(data_ws, min_col=category_col, min_row=2, max_row=max_row)

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.legend = None
    chart.varyColors = False
    chart.display_blanks = "gap"

    if chart.series:
        style_line_series(chart.series[0], color, marker_size)

    ws.add_chart(chart, anchor)


def add_column_chart(
    ws,
    data_ws,
    title,
    min_col,
    max_row,
    category_col,
    anchor,
    values,
    date_format,
    major_time_unit,
    color,
    y_major_unit,
    gap_width=15,
    x_min_date=None,
    x_max_date=None,
    height=9,
    width=18.5,
):
    chart = BarChart()
    chart.type = "col"
    chart.grouping = "clustered"
    chart.overlap = 0
    chart.gapWidth = gap_width
    chart.title = title
    chart.height = height
    chart.width = width

    chart.x_axis = DateAxis(axId=500, crossAx=100)
    chart.x_axis.number_format = date_format
    chart.x_axis.majorTimeUnit = major_time_unit
    chart.x_axis.tickLblPos = "nextTo"
    chart.x_axis.tickLblSkip = 1
    chart.x_axis.tickMarkSkip = 1
    chart.x_axis.tickMark = "out"
    chart.x_axis.delete = False

    if x_min_date is not None:
        chart.x_axis.scaling.min = to_excel(x_min_date)
    if x_max_date is not None:
        chart.x_axis.scaling.max = to_excel(x_max_date)

    chart.y_axis.axId = 100
    chart.y_axis.crossAx = 500
    chart.y_axis.numFmt = "#,##0"
    chart.y_axis.tickLblPos = "nextTo"
    chart.y_axis.tickMark = "out"
    chart.y_axis.delete = False
    chart.y_axis.scaling.min = 0
    chart.y_axis.scaling.max, chart.y_axis.majorUnit = axis_scale(
        values,
        y_major_unit,
    )
    chart.y_axis.majorGridlines = ChartLines()

    data = Reference(data_ws, min_col=min_col, min_row=1, max_row=max_row)
    cats = Reference(data_ws, min_col=category_col, min_row=2, max_row=max_row)

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.legend = None
    chart.varyColors = False
    chart.display_blanks = "gap"

    if chart.series:
        style_bar_series(chart.series[0], color)

    ws.add_chart(chart, anchor)


def add_category_line_chart(
    ws,
    title,
    data_col,
    max_row,
    header_row,
    first_data_row,
    anchor,
    x_axis_title,
    y_axis_title,
    color,
):
    if max_row < first_data_row:
        return

    chart = LineChart()
    chart.title = title
    chart.height = 9
    chart.width = 18.5
    chart.x_axis.tickLblPos = "nextTo"
    chart.x_axis.tickMark = "out"
    chart.x_axis.delete = False
    chart.y_axis.numFmt = "#,##0"
    chart.y_axis.tickLblPos = "nextTo"
    chart.y_axis.tickMark = "out"
    chart.y_axis.delete = False
    chart.y_axis.scaling.min = 0
    values = [
        ws.cell(row, data_col).value
        for row in range(first_data_row, max_row + 1)
    ]
    chart.y_axis.scaling.max, chart.y_axis.majorUnit = nice_axis_scale(values)
    chart.y_axis.majorGridlines = ChartLines()

    data = Reference(ws, min_col=data_col, min_row=header_row, max_row=max_row)
    cats = Reference(ws, min_col=1, min_row=first_data_row, max_row=max_row)

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.legend = None
    chart.varyColors = False

    if chart.series:
        style_line_series(chart.series[0], color)

    ws.add_chart(chart, anchor)


def make_calendar(ws, activities, year, start_row=4, start_col=10):
    ws.merge_cells(
        start_row=start_row,
        start_column=start_col,
        end_row=start_row,
        end_column=start_col + 7,
    )
    ws.cell(start_row, start_col, f"Calendrier des activites {year}")
    style_section_cell(ws.cell(start_row, start_col))
    ws.cell(start_row, start_col).alignment = Alignment(horizontal="center")

    day_categories = calendar_day_categories(activities, year)
    category_colors = {
        "running": RUNNING_COLOR,
        "trail": TRAIL_COLOR,
        "walk": WALK_COLOR,
        "bike": BIKE_COLOR,
        "other": OTHER_COLOR,
    }
    month_names = [
        "Janvier", "Fevrier", "Mars", "Avril", "Mai", "Juin",
        "Juillet", "Aout", "Septembre", "Octobre", "Novembre", "Decembre",
    ]

    header_row = start_row + 1
    for offset, label in enumerate(["L", "M", "M", "J", "V", "S", "D"]):
        cell = ws.cell(header_row, start_col + offset)
        cell.value = label
        cell.alignment = Alignment(horizontal="center")
        cell.font = Font(bold=True, size=9)

    ws.cell(header_row, start_col + 7, "Mois")
    ws.cell(header_row, start_col + 7).font = Font(bold=True, size=9)

    current = week_start(date(year, 1, 1))
    end = week_start(date(year, 12, 31))
    row = header_row + 1

    while current <= end:
        week_days = [current + timedelta(days=offset) for offset in range(7)]
        visible_months = []

        for offset, current_day in enumerate(week_days):
            cell = ws.cell(row, start_col + offset)
            cell.alignment = Alignment(horizontal="center")
            cell.border = BORDER
            cell.font = Font(size=10)

            if current_day.year != year:
                cell.value = ""
                cell.fill = PatternFill("solid", fgColor="F9FAFB")
                continue

            cell.value = current_day.day
            category = day_categories.get(current_day)

            if category:
                cell.fill = PatternFill("solid", fgColor=category_colors[category])
                cell.font = Font(color="FFFFFF", bold=True, size=10)
            else:
                cell.fill = PatternFill("solid", fgColor="FFFFFF")

            month_name = month_names[current_day.month - 1]
            if month_name not in visible_months:
                visible_months.append(month_name)

        month_cell = ws.cell(row, start_col + 7)
        month_cell.value = " / ".join(visible_months)
        month_cell.alignment = Alignment(horizontal="left")
        month_cell.font = Font(size=9, color="4B5563")
        ws.row_dimensions[row].height = 22
        row += 1
        current += timedelta(days=7)

    for col in range(start_col, start_col + 7):
        ws.column_dimensions[get_column_letter(col)].width = 6.5
    ws.column_dimensions[get_column_letter(start_col + 7)].width = 20
    ws.row_dimensions[start_row].height = 24
    ws.row_dimensions[header_row].height = 20


def make_dashboard(wb, activities):
    ws = wb["Dashboard"]
    setup_sheet(ws, "A4")

    today = datetime.now().date()
    year = today.year
    sport_activities = [a for a in activities if is_running_or_trail(a)]

    write_title(
        ws,
        "Sport Dashboard",
        f"Annee en cours : {year} - course a pied et trail",
    )

    headers = [
        "Periode",
        "Sorties",
        "Course",
        "Trail",
        "Distance km",
        "Duree h",
        "D+",
        "Allure moy.",
    ]
    ws.append([])
    ws.append(headers)
    style_header(ws[4])

    for label, start_day, end_day in current_periods(today):
        stats = period_stats(activities, start_day, end_day)
        ws.append([
            label,
            stats["activities"],
            stats["running"],
            stats["trail"],
            stats["distance_km"],
            stats["duration_h"],
            stats["elevation_gain_m"],
            stats["pace"],
        ])

    for row in ws.iter_rows(min_row=5, max_row=7, min_col=1, max_col=8):
        for cell in row:
            cell.fill = CARD_FILL
            cell.border = BORDER

    format_columns(ws, {"E": "0.0", "F": "0.0", "G": "#,##0", "H": "m:ss"}, 5)

    day_rows, total_days = yearly_day_breakdown(activities, year, today)
    ws.merge_cells("S4:U4")
    ws["S4"] = f"Jours en {year} - {total_days} jours"
    style_section_cell(ws["S4"])
    ws["S4"].alignment = Alignment(horizontal="center")

    for column, header in enumerate(["Type", "Jours", "Pourcentage"], start=19):
        ws.cell(5, column, header)
    style_header(ws[5][18:21])

    for row_index, (label, count, percentage, color) in enumerate(day_rows, start=6):
        font_color = "111827" if label == "Sans activite" else "FFFFFF"

        for column, value in enumerate([label, count, percentage], start=19):
            cell = ws.cell(row_index, column, value)
            cell.fill = PatternFill("solid", fgColor=color)
            cell.font = Font(color=font_color, bold=True, size=9)
            cell.border = BORDER
            cell.alignment = Alignment(
                horizontal="left" if column == 19 else "center"
            )

        ws.cell(row_index, 21).number_format = "0.0%"

    year_rows = daily_year_data(activities, year, today)
    weekly_current = weekly_year_data(activities, year, today)
    monthly_current = monthly_year_data(activities, year, today)
    year_elevation_rows = daily_year_elevation_data(activities, year, today)
    weekly_elevation_current = weekly_year_elevation_data(activities, year, today)
    year_start = date(year, 1, 1)
    year_end = date(year, 12, 31)

    data_ws = wb.create_sheet("_Data")
    write_chart_data(
        data_ws,
        year_rows,
        weekly_current,
        monthly_current,
        year_elevation_rows,
        weekly_elevation_current,
    )

    year_total_km = next(
        (value for _, value in reversed(year_rows) if value is not None),
        0,
    )
    vo2max = latest_non_null(sport_activities, "vo2max")

    ws["A9"] = f"Total {year}"
    ws["B9"] = year_total_km
    ws["C9"] = "km"
    ws["E9"] = "VO2max actuel"
    ws["F9"] = vo2max

    for cell in ["A9", "B9", "C9", "E9", "F9"]:
        ws[cell].fill = ACCENT_FILL if cell in {"A9", "E9"} else CARD_FILL
        ws[cell].border = BORDER
        ws[cell].font = Font(bold=True)

    ws["B9"].number_format = "0.0"
    ws["F9"].number_format = "0.0"

    add_line_chart(
        ws,
        data_ws,
        f"Cumul kilometrique {year} (km)",
        min_col=2,
        max_row=len(year_rows) + 1,
        category_col=1,
        anchor="A11",
        values=[row[1] for row in year_rows],
        x_axis_title=None,
        y_axis_title=None,
        date_format="mmm",
        major_time_unit="months",
        color=CHART_BLUE,
        y_major_unit=250,
        marker_size=2,
        x_min_date=year_start,
        x_max_date=year_end,
        height=8,
        width=17.2,
    )
    add_column_chart(
        ws,
        data_ws,
        f"Distance hebdomadaire {year} (km)",
        min_col=5,
        max_row=len(weekly_current) + 1,
        category_col=4,
        anchor="A22",
        values=[row[1] for row in weekly_current],
        date_format="dd mmm",
        major_time_unit="months",
        color=CHART_GREEN,
        y_major_unit=5,
        gap_width=5,
        x_min_date=year_start,
        x_max_date=year_end,
        height=8,
        width=17.2,
    )
    add_column_chart(
        ws,
        data_ws,
        f"Distance mensuelle {year} (km)",
        min_col=8,
        max_row=len(monthly_current) + 1,
        category_col=7,
        anchor="A33",
        values=[row[1] for row in monthly_current],
        date_format="mmm",
        major_time_unit="months",
        color=CHART_ORANGE,
        y_major_unit=20,
        gap_width=20,
        x_min_date=year_start,
        x_max_date=year_end,
        height=8,
        width=17.2,
    )
    add_line_chart(
        ws,
        data_ws,
        f"Cumul D+ {year} (m)",
        min_col=11,
        max_row=len(year_elevation_rows) + 1,
        category_col=10,
        anchor="A44",
        values=[row[1] for row in year_elevation_rows],
        x_axis_title=None,
        y_axis_title=None,
        date_format="mmm",
        major_time_unit="months",
        color=CHART_TEAL,
        y_major_unit=1000,
        marker_size=2,
        x_min_date=year_start,
        x_max_date=year_end,
        height=8,
        width=17.2,
    )
    add_column_chart(
        ws,
        data_ws,
        f"D+ hebdomadaire {year} (m)",
        min_col=14,
        max_row=len(weekly_elevation_current) + 1,
        category_col=13,
        anchor="A55",
        values=[row[1] for row in weekly_elevation_current],
        date_format="dd mmm",
        major_time_unit="months",
        color=CHART_TEAL,
        y_major_unit=250,
        gap_width=5,
        x_min_date=year_start,
        x_max_date=year_end,
        height=8,
        width=17.2,
    )

    autosize(ws)
    make_calendar(ws, activities, year, start_row=4, start_col=10)
    ws.column_dimensions["A"].width = 17
    ws.column_dimensions["H"].width = 11.5
    ws.column_dimensions["I"].width = 2
    ws.column_dimensions["R"].width = 3
    ws.column_dimensions["S"].width = 18
    ws.column_dimensions["T"].width = 9
    ws.column_dimensions["U"].width = 13

    for row in ws.iter_rows(
        min_row=1,
        max_row=ws.max_row,
        min_col=1,
        max_col=ws.max_column,
    ):
        for cell in row:
            if cell.value is None:
                continue

            alignment = copy(cell.alignment)
            alignment.vertical = "center"
            cell.alignment = alignment


def format_activities(ws):
    formats = {
        "A": "yyyy-mm-dd",
        "B": "hh:mm",
        "F": "0.00",
        "G": "[h]:mm:ss",
        "H": "m:ss",
        "I": "#,##0",
        "J": "#,##0",
        "K": "0",
        "L": "0",
        "M": "0.0",
        "N": "0",
        "O": "0",
        "P": "0",
        "Q": "0.0",
        "R": "0.0",
        "S": "0.0",
        "T": "#,##0",
    }
    format_columns(ws, formats)


def format_weeks(ws):
    formats = {
        "C": "yyyy-mm-dd",
        "H": "0.00",
        "I": "0.0",
        "J": "m:ss",
        "K": "#,##0",
        "L": "#,##0",
        "M": "0.00",
        "N": "0.00",
        "O": "#,##0",
        "P": "0",
        "Q": "0",
    }
    format_columns(ws, formats)


def highlight_week_maxima(ws):
    if ws.max_row < 4:
        return

    # Keep identifiers and the weighted FC/power context neutral.
    for column in ["E", "H", "I", "K", "L", "M", "N", "O"]:
        cells = [
            ws[f"{column}{row}"]
            for row in range(4, ws.max_row + 1)
            if isinstance(ws[f"{column}{row}"].value, (int, float))
        ]

        if not cells:
            continue

        maximum = max(cell.value for cell in cells)

        for cell in cells:
            if cell.value == maximum:
                cell.fill = MAXIMUM_FILL
                cell.font = Font(bold=True, color="854D0E")


def make_records_sheet(ws, activities, weeks):
    setup_sheet(ws)
    write_title(ws, "Records", "Course a pied et trail uniquement")

    sections = [
        (
            "Records route",
            ["Distance", "Temps", "Allure", "Date", "Activite", "Garmin ID"],
            route_record_rows(activities),
            "TableRecordsRoute",
        ),
        (
            "Records trail",
            ["Record", "Valeur", "Unite", "Date", "Activite", "Garmin ID"],
            trail_record_rows(activities),
            "TableRecordsTrail",
        ),
        (
            "General",
            ["Record", "Valeur", "Unite", "Date / semaine", "Reference", "Garmin ID"],
            general_record_rows(activities, weeks),
            "TableRecordsGeneral",
        ),
    ]

    current_row = 4

    for title, headers, rows, table_name in sections:
        ws.cell(current_row, 1, title)
        style_section_cell(ws.cell(current_row, 1))
        current_row += 1

        for col_index, header in enumerate(headers, start=1):
            ws.cell(current_row, col_index, header)
        style_header(ws[current_row])

        first_row = current_row
        current_row += 1

        for row in rows:
            for col_index, value in enumerate(row, start=1):
                ws.cell(current_row, col_index, value)
            current_row += 1

        add_table(ws, table_name, first_row, current_row - 1, len(headers))
        current_row += 2

    for row in range(6, 10):
        ws.cell(row, 2).number_format = "[h]:mm:ss"
        ws.cell(row, 3).number_format = "m:ss"
        ws.cell(row, 4).number_format = "yyyy-mm-dd"

    for row in range(13, 16):
        ws.cell(row, 4).number_format = "yyyy-mm-dd"

    autosize(ws)


def make_all_sheet(ws, activities):
    setup_sheet(ws)
    write_title(ws, "All", "Comparaison annuelle - course a pied et trail")

    rows = yearly_rows(activities)
    headers = [
        "Annee",
        "Nb seances",
        "Course",
        "Trail",
        "Distance km",
        "Duree h",
        "D+",
        "D-",
        "Sortie max km",
        "FC moy. ponderee",
        "Puissance max",
        "VO2max max",
    ]

    ws.append([])
    ws.append(headers)
    style_header(ws[3])

    for row in rows:
        ws.append(row)

    if rows:
        add_table(ws, "TableAllYears", 4, ws.max_row, len(headers))
        max_row = ws.max_row
        chart_start = max_row + 3

        add_category_line_chart(
            ws,
            "Distance par an",
            data_col=5,
            max_row=max_row,
            header_row=4,
            first_data_row=5,
            anchor=f"A{chart_start}",
            x_axis_title="Annee",
            y_axis_title="Distance (km)",
            color=CHART_BLUE,
        )
        add_category_line_chart(
            ws,
            "D+ par an",
            data_col=7,
            max_row=max_row,
            header_row=4,
            first_data_row=5,
            anchor=f"A{chart_start + 20}",
            x_axis_title="Annee",
            y_axis_title="Denivele positif (m)",
            color=CHART_GREEN,
        )
        add_category_line_chart(
            ws,
            "Duree par an",
            data_col=6,
            max_row=max_row,
            header_row=4,
            first_data_row=5,
            anchor=f"A{chart_start + 40}",
            x_axis_title="Annee",
            y_axis_title="Duree (h)",
            color=CHART_ORANGE,
        )

    format_columns(
        ws,
        {
            "E": "0.00",
            "F": "0.0",
            "G": "#,##0",
            "H": "#,##0",
            "I": "0.00",
            "J": "0",
            "K": "0",
            "L": "0.0",
        },
    )
    autosize(ws)


def build_workbook(activities):
    wb = Workbook()
    wb.active.title = "Dashboard"
    ws_weeks = wb.create_sheet("Semaines")
    ws_activities = wb.create_sheet("Activites")
    ws_records = wb.create_sheet("Records")
    ws_all = wb.create_sheet("All")

    weeks = weekly_rows(activities)

    make_dashboard(wb, activities)

    make_table_sheet(
        ws_weeks,
        "Semaines",
        [
            "Annee",
            "Semaine",
            "Debut semaine",
            "Label",
            "Nb seances",
            "Course",
            "Trail",
            "Distance km",
            "Duree h",
            "Allure moy.",
            "D+",
            "D-",
            "Distance moy./seance",
            "Sortie max km",
            "D+ max sortie",
            "FC moy. ponderee",
            "Puissance moy. ponderee",
        ],
        weeks,
        "TableSemaines",
        format_weeks,
    )
    highlight_week_maxima(ws_weeks)

    make_table_sheet(
        ws_activities,
        "Activites",
        [
            "Date",
            "Heure",
            "Nom",
            "Categorie",
            "Type Garmin",
            "Distance km",
            "Duree",
            "Allure",
            "D+",
            "D-",
            "FC moy.",
            "FC max.",
            "Cadence",
            "Puissance moy.",
            "Puissance norm.",
            "Puissance max.",
            "TE aerobie",
            "TE anaerobie",
            "VO2max",
            "Calories",
            "Garmin ID",
        ],
        activity_rows(activities),
        "TableActivites",
        format_activities,
    )

    make_records_sheet(ws_records, activities, weeks)
    make_all_sheet(ws_all, activities)

    wb.calculation.fullCalcOnLoad = True
    return wb


def main():
    activities = load_activities()
    EXPORT_PATH.parent.mkdir(parents=True, exist_ok=True)

    wb = build_workbook(activities)
    wb.save(EXPORT_PATH)

    print(f"Excel genere : {EXPORT_PATH}")
    print(f"Activites chargees : {len(activities)}")
    print(f"Taille : {EXPORT_PATH.stat().st_size} octets")


if __name__ == "__main__":
    main()
